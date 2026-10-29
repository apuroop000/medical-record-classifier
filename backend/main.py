import os
import shutil
import joblib
import pandas as pd
import fitz  # PyMuPDF
from PIL import Image, ImageOps
import pytesseract
import numpy as np
from datetime import datetime
from fastapi import FastAPI, UploadFile, File, Form, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from typing import List, Dict

try:
    import docx
except ImportError:
    docx = None

app = FastAPI(title="Intelligent Medical Record Router API")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

rpa_queue: List[Dict] = []
job_counter = 1
AUDIT_EXCEL_PATH = os.path.join("data", "audit_log.xlsx")

MODEL_PATH = "classifier/saved_model/model.pkl"
model_bundle = None

if os.path.exists(MODEL_PATH):
    model_bundle = joblib.load(MODEL_PATH)
    print("✅ Model bundle loaded successfully in FastAPI backend!")
else:
    print("⚠️ Warning: Model bundle not found. Falling back to keyword + image heuristics only.")

DEPARTMENT_MAP = {
    "Laboratory / Blood Test Report": "Laboratory",
    "X-Ray Report": "Radiology",
    "ECG Report": "Cardiology",
    "Prescription": "Pharmacy",
    "Discharge Summary": "Medical Records",
    "Unclassified – Needs Review": "Unassigned",
}

SUPPORTED_IMAGE_EXT = [".png", ".jpg", ".jpeg", ".bmp", ".tiff", ".tif"]
SUPPORTED_DOC_EXT = [".pdf", ".docx"]


def log_to_excel(job_data: dict):
    os.makedirs("data", exist_ok=True)
    row = {
        "Job ID": job_data["id"],
        "Timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "Patient Name": job_data["patient_name"],
        "Patient ID": job_data["patient_id"],
        "Filename": job_data["filename"],
        "Document Type": job_data["document_type"],
        "Routed Department": job_data["department"],
        "Confidence": job_data["confidence"],
        "Method": job_data["classification_method"],
        "Status": job_data["status"]
    }
    
    if os.path.exists(AUDIT_EXCEL_PATH):
        try:
            df_existing = pd.read_excel(AUDIT_EXCEL_PATH)
            df_new = pd.concat([df_existing, pd.DataFrame([row])], ignore_index=True)
        except Exception:
            df_new = pd.DataFrame([row])
    else:
        df_new = pd.DataFrame([row])
        
    df_new.to_excel(AUDIT_EXCEL_PATH, index=False)


def preprocess_image_for_ocr(image: Image.Image) -> Image.Image:
    image = ImageOps.exif_transpose(image)
    image = image.convert("L")
    if image.width < 1500:
        scale = 1500 / image.width
        image = image.resize((int(image.width * scale), int(image.height * scale)))
    image = ImageOps.autocontrast(image)
    image = image.point(lambda x: 0 if x < 150 else 255, "1")
    return image


def looks_like_radiograph(image: Image.Image) -> bool:
    try:
        rgb = image.convert("RGB")
        arr = np.array(rgb).astype(int)
        r, g, b = arr[:, :, 0], arr[:, :, 1], arr[:, :, 2]
        color_spread = np.mean(np.abs(r - g) + np.abs(g - b) + np.abs(r - b))
        return color_spread < 25
    except Exception:
        return False


def extract_text_from_pdf(file_path: str) -> str:
    text = ""
    doc = fitz.open(file_path)
    for page in doc:
        text += page.get_text() + "\n"
    doc.close()

    if len(text.strip()) < 20:
        text = ""
        doc = fitz.open(file_path)
        for page in doc:
            pix = page.get_pixmap(dpi=300)
            img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
            processed = preprocess_image_for_ocr(img)
            text += pytesseract.image_to_string(processed, config="--psm 6") + "\n"
        doc.close()
    return text


def extract_text_from_docx(file_path: str) -> str:
    if docx is None:
        return ""
    text_parts = []
    document = docx.Document(file_path)
    for para in document.paragraphs:
        text_parts.append(para.text)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                text_parts.append(cell.text)
    return "\n".join(text_parts)


def extract_text_from_image(file_path: str) -> str:
    image = Image.open(file_path)
    processed = preprocess_image_for_ocr(image)
    return pytesseract.image_to_string(processed, config="--psm 6")


def extract_text_from_file(file_path: str, ext: str) -> str:
    text = ""
    try:
        if ext == ".pdf":
            text = extract_text_from_pdf(file_path)
        elif ext == ".docx":
            text = extract_text_from_docx(file_path)
        elif ext in SUPPORTED_IMAGE_EXT:
            text = extract_text_from_image(file_path)
    except Exception as e:
        print(f"❌ Error extracting text from {file_path}: {e}")
    return text


KEYWORD_RULES = [
    ("ECG Report", [
        "ecg report", "electrocardiogram", "ekg", "cardiac rhythm",
        "cardiology department", "heart rate strip", "ecg test"
    ]),
    ("X-Ray Report", [
        "x-ray", "xray", "radiology report", "radiograph", "fracture",
        "bone density", "chest x ray", "ct scan", "mri scan"
    ]),
    ("Discharge Summary", [
        "discharge summary", "hospital admission", "discharged on",
        "date of admission", "date of discharge", "course in hospital"
    ]),
    ("Prescription", [
        "prescription", "rx", "tablet", "capsule", "dosage", "mg twice daily",
        "opd-f", "dr.", "signature of doctor", "take after food", "yashoda hospital"
    ]),
    ("Laboratory / Blood Test Report", [
        "complete blood count", "cbc", "blood test report", "pathology lab",
        "serum", "hemoglobin", "lab report", "specimen collected",
        "reference range"
    ]),
]


def classify_by_keywords(combined_content: str):
    for doc_type, keywords in KEYWORD_RULES:
        if any(kw in combined_content for kw in keywords):
            return doc_type, 0.95
    return None, None


def classify_with_ml_model(doc_text: str, fallback_text: str):
    if not model_bundle:
        return None, None
    vectorizer = model_bundle["vectorizer"]
    model = model_bundle["model"]
    le = model_bundle["label_encoder"]

    input_text = doc_text if doc_text.strip() else fallback_text
    X_vec = vectorizer.transform([input_text])

    if hasattr(model, "predict_proba"):
        probs = model.predict_proba(X_vec)
        pred_idx = int(np.argmax(probs, axis=1)[0])
        confidence = float(np.max(probs))
    else:
        pred_idx = int(model.predict(X_vec)[0])
        confidence = 0.6

    doc_type = le.inverse_transform([pred_idx])[0]
    if confidence < 0.55:
        return None, None
    return doc_type, confidence


def classify_document(file_path: str, filename: str, ext: str, doc_text: str):
    combined_content = (filename + " " + doc_text).lower()

    doc_type, confidence = classify_by_keywords(combined_content)
    if doc_type:
        return doc_type, confidence, "keyword"

    if ext in SUPPORTED_IMAGE_EXT:
        try:
            image = Image.open(file_path)
            if looks_like_radiograph(image):
                return "X-Ray Report", 0.80, "image-heuristic"
        except Exception:
            pass

    doc_type, confidence = classify_with_ml_model(doc_text, filename)
    if doc_type:
        return doc_type, confidence, "ml-model"

    return "Unclassified – Needs Review", 0.0, "none"


@app.post("/api/process-document")
async def process_document(
    file: UploadFile = File(...),
    patient_name: str = Form(...),
    patient_id: str = Form(...)
):
    global job_counter
    try:
        patient_folder = os.path.join("data", "patient_records", patient_name.strip())
        os.makedirs(patient_folder, exist_ok=True)

        file_path = os.path.join(patient_folder, file.filename)
        with open(file_path, "wb") as buffer:
            shutil.copyfileobj(file.file, buffer)

        _, ext = os.path.splitext(file.filename.lower())

        if ext not in SUPPORTED_IMAGE_EXT and ext not in SUPPORTED_DOC_EXT:
            raise HTTPException(status_code=400, detail=f"Unsupported file type '{ext}'.")

        doc_text = extract_text_from_file(file_path, ext)
        doc_type, confidence, method = classify_document(file_path, file.filename, ext, doc_text)
        department = DEPARTMENT_MAP.get(doc_type, "General")

        job = {
            "id": job_counter,
            "filename": file.filename,
            "file_path": file_path,
            "patient_name": patient_name.strip(),
            "patient_id": patient_id,
            "document_type": doc_type,
            "department": department,
            "confidence": confidence,
            "classification_method": method,
            "status": "pending"
        }
        
        rpa_queue.append(job)
        log_to_excel(job)
        job_counter += 1

        return {
            "status": "success",
            "filename": file.filename,
            "file_path": file_path,
            "patient_name": patient_name.strip(),
            "document_type": doc_type,
            "routed_department": department,
            "confidence": confidence,
            "classification_method": method
        }
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.get("/api/rpa/pending")
async def get_pending_jobs():
    pending = [j for j in rpa_queue if j["status"] == "pending"]
    return {"jobs": pending}


@app.post("/api/rpa/run-bot")
async def run_rpa_bot():
    """Fixed RPA Bot worker that loops through all pending queue items, updates status, and logs completion."""
    processed_count = 0
    for job in rpa_queue:
        if job["status"].lower() == "pending":
            job["status"] = "Completed & Routed"
            log_to_excel(job)
            processed_count += 1
            
    return {
        "status": "success",
        "message": f"🤖 RPA Bot successfully executed and processed {processed_count} pending document records!",
        "processed_count": processed_count
    }


@app.post("/api/rpa/update/{job_id}")
async def update_job_status(job_id: int, status: str = Form(...)):
    for job in rpa_queue:
        if job["id"] == job_id:
            job["status"] = status
            log_to_excel(job)
            return {"status": "success", "job_id": job_id, "new_status": status}
    raise HTTPException(status_code=404, detail="Job not found")
